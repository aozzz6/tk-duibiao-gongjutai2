from __future__ import annotations
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import json
import os
from pathlib import Path
WEBAPP_DIR = Path(__file__).resolve().parent
PIPELINE_DIR = Path(os.environ.get('WB_PIPELINE_DIR', os.environ.__getitem__('WB_PIPELINE_DIR'))).resolve()
VENV_PYTHON = Path(os.environ.get('WB_PYTHON', str(PIPELINE_DIR / '.venv' / 'Scripts' / 'python.exe')))
PROXY = os.environ.get('WB_PROXY', 'http://127.0.0.1:7897')
LEDGER_DIR = Path(os.environ.get('WB_LEDGER_DIR', os.environ.__getitem__('WB_LEDGER_DIR')))
AUTO_TRANSLATE_ENABLED = os.environ.get('WB_AUTO_TRANSLATE', '1') not in ('0', 'false', 'False')
AUTO_TRANSLATE_MODEL = os.environ.get('WB_AUTO_TRANSLATE_MODEL', '')
OUTPUT_DIR = PIPELINE_DIR / 'output'
REPORTS_DIR = PIPELINE_DIR / 'reports'
SCRIPTS_DIR = PIPELINE_DIR / 'scripts'
ACCOUNTS_JSON = PIPELINE_DIR / 'accounts.json'
STATE_FILE = PIPELINE_DIR / 'pipeline_state.json'
GATE_REPORT = PIPELINE_DIR / 'gate_report.json'
RUN_LOG = PIPELINE_DIR / 'pipeline_run.log'
COOKIE_FILE = PIPELINE_DIR / 'cc_cookie.txt'
DATA_DIR = WEBAPP_DIR / 'data'
RUNS_JSONL = DATA_DIR / 'runs.jsonl'
SCHEDULE_JSON = DATA_DIR / 'schedule.json'
SCHEDULE_RETRY = int(os.environ.get('WB_SCHEDULE_RETRY', '2'))
SCHEDULE_TIMEOUT = int(os.environ.get('WB_SCHEDULE_TIMEOUT', '5400'))
AUTOMATIONS_JSON = DATA_DIR / 'automations.json'
for _d in (DATA_DIR,):
    _d.mkdir(parents=True, exist_ok=True)
PIPELINE_STEPS = ['s0_check', 's1_download', 's2_scrape', 's3_4_sheet', 's5_transcribe', 's5b_products', 's6_translate', 's7_backfill', 's9_report']
STEP_LABELS = {'s0_check': '前置校验', 's1_download': '音频下载', 's2_scrape': '页面抓取', 's3_4_sheet': '建表录入', 's5_transcribe': '转写', 's5b_products': '商品提取', 's6_translate': '翻译', 's7_backfill': '台账回填', 's9_report': '报表'}
AGENT_STEPS = set()
LOCAL_STEPS = [s for s in PIPELINE_STEPS if s != 's6_translate']
DAILY_STEPS = ['s0_check', 's1_download', 's2_scrape', 's5_transcribe']
RUN_EXTRA_STEPS = ['s5b_products']
SHEET_PREFIX = '对标记录-'
NOTE_SHEET_KEY = '说明'
EXPECTED_COLS = ['记录日期', '对标账号', '视频发布时间', '视频标题', '话题标签', '视频链接/ID', '商品链接', '播放量', '点赞数', '评论数', '转发数', '收藏数', '音频', '翻译']
IMAGE_EXPECTED_COLS = ['记录日期', '对标账号', '图文发布时间', '图文标题', '话题标签', '视频链接/ID', '商品链接', '播放量', '点赞数', '评论数', '转发数', '收藏数']
COL_ALIASES = {'记录日期': ['记录日期', '日期', '记录时间', 'date'], '对标账号': ['对标账号', '账号', 'account'], '视频发布时间': ['视频发布时间', '发布时间', 'publish'], '视频链接/ID': ['视频链接/ID', '视频链接ID', '视频链接', '链接/ID', '视频ID', 'video_id'], '视频出单量': ['视频出单量', '出单量', 'GMV', 'gmv'], '商品链接': ['商品链接', '商品', 'product', 'bio_link'], '音频': ['音频', '备注', '转写'], '翻译': ['翻译', '备注中文翻译', '中文翻译']}

def read_json(path, default=None):
    try:
        return json.loads(Path(path).read_text(encoding='utf-8'))
    except Exception:
        return default

def python_exe() -> str:
    if VENV_PYTHON.exists():
        return str(VENV_PYTHON)
    return 'python'

def load_accounts(include_inactive: bool=False) -> dict:
    """读 accounts.json。

    include_inactive=False（缺省）时过滤掉 active=False 的账号——
    让停用账号自动退出流水线执行链，而历史产物索引仍保留在文件里。
    注：工作台 UI 需要看到全部账号（含停用），调用时传 True。
    """
    cfg = read_json(ACCOUNTS_JSON, {}) or {}
    if include_inactive:
        return cfg
    accs = cfg.get('accounts')
    if isinstance(accs, list):
        cfg = dict(cfg)
        cfg['accounts'] = [a for a in accs if not isinstance(a, dict) or a.get('active', True)]
    return cfg

def load_accounts_raw() -> dict:
    """不过滤，返回文件原样（含停用账号）。"""
    return read_json(ACCOUNTS_JSON, {}) or {}
ACCOUNT_RE = __import__('re').compile('^[a-z0-9_.]+$')
_SHEET_BAD = set('[]:*?/\\')

def validate_account_name(name: str) -> str:
    """校验账号名（同时是 Excel sheet 名后缀）。返回规范化后的名字，非法则抛 ValueError。"""
    if not isinstance(name, str):
        raise ValueError('账号名必须是字符串')
    name = (name or '').strip()
    if not name:
        raise ValueError('账号名不能为空')
    if not ACCOUNT_RE.match(name):
        raise ValueError('账号名只能含小写字母、数字、下划线、点（与 TikTok 用户名一致）')
    if len(SHEET_PREFIX + name) > 31:
        raise ValueError(f"账号名过长：{'对标记录-' + name} 超过 Excel sheet 名 31 字符上限")
    if _SHEET_BAD & set(name):
        raise ValueError('账号名不能含 []:*?/\\ 等 Excel sheet 非法字符')
    return name

def default_account(name: str, lang: str='en') -> dict:
    """按主流水线约定生成新账号的缺省配置（out_dir 统一为 pipeline/output）。"""
    return {'account': name, 'lang': lang or 'en', 'out_dir': str(OUTPUT_DIR), 'audio_dir': str(OUTPUT_DIR / f'{name}_audio')}
REGION_LANG_MAP = {'us': 'en', 'uk': 'en', 'eu': 'en', 'de': 'de', 'fr': 'fr', 'es': 'es', 'it': 'it', 'cn': 'zh', 'cn-hk': 'zh', 'cn-tw': 'zh', 'jp': 'jp', 'kr': 'kr'}
SUPPORTED_LANGS = {'en', 'de', 'fr', 'es', 'it', 'zh', 'jp', 'kr'}
SUPPORTED_REGIONS = set(REGION_LANG_MAP) | {''}

def validate_account_fields(fields: dict) -> dict:
    """校验账号字段，避免 API 直接写入与界面 schema 不一致的数据。"""
    fields = fields or {}
    if not isinstance(fields, dict):
        raise ValueError('账号字段必须是对象')
    if 'lang' in fields and fields['lang'] not in (None, ''):
        if not isinstance(fields['lang'], str) or fields['lang'] not in SUPPORTED_LANGS:
            raise ValueError(f"不支持的转写语言: {fields['lang']}")
    if 'region' in fields and fields['region'] not in (None, ''):
        if not isinstance(fields['region'], str) or fields['region'].strip().lower() not in SUPPORTED_REGIONS:
            raise ValueError(f"不支持的面向地区: {fields['region']}")
        fields = dict(fields)
        fields['region'] = fields['region'].strip().lower()
    for key in ('nickname', 'note', 'out_dir', 'audio_dir'):
        if key in fields and fields[key] is not None and (not isinstance(fields[key], str)):
            raise ValueError(f'{key} 必须是字符串')
    if 'tags' in fields and fields['tags'] is not None:
        tags = fields['tags']
        if isinstance(tags, str):
            fields = dict(fields)
            fields['tags'] = [t.strip() for t in tags.split(',') if t.strip()]
        elif not isinstance(tags, list) or any((not isinstance(t, str) for t in tags)):
            raise ValueError('tags 必须是字符串或字符串数组')
    if 'active' in fields and (not isinstance(fields['active'], bool)):
        raise ValueError('active 必须是布尔值')
    return fields

def guess_lang(region: str | None, explicit_lang: str | None=None) -> str:
    """按 region 推断转写语言；显式传入的 lang 优先。

    返回 (lang, source) 由调用方处理；这里只回 lang。
    - 有显式 lang 且合法 → 用显式（最高优先）
    - 无显式 → 按 region 查 REGION_LANG_MAP，缺省 en
    """
    if explicit_lang and explicit_lang in ('en', 'de', 'fr', 'es', 'it', 'zh', 'jp', 'kr'):
        return explicit_lang
    r = (region or '').strip().lower()
    if r in REGION_LANG_MAP:
        return REGION_LANG_MAP[r]
    return 'en'

def _write_accounts(cfg: dict) -> None:
    """原子写：先写 .tmp 再 os.replace，避免流水线读到半个文件。"""
    tmp = ACCOUNTS_JSON.with_suffix(ACCOUNTS_JSON.suffix + '.tmp')
    text = json.dumps(cfg, ensure_ascii=False, indent=2) + '\n'
    tmp.write_text(text, encoding='utf-8')
    os.replace(tmp, ACCOUNTS_JSON)
DEFAULT_ACCOUNT_FIELDS = {'nickname': '', 'active': True, 'tags': [], 'note': ''}
HEALTH_LABELS = {'ok': ('正常', 'ok'), 'stale': ('陈旧', 'warn'), 'partial': ('缺件', 'warn'), 'config_only': ('仅配置，无产物', 'error'), 'off': ('已停用', 'gray')}
STALE_DAYS = 7

def _account_norm(acc: dict) -> dict:
    """补齐可选字段（不改原对象）。关键：active 缺省 True，
    否则现有 accounts.json 无该字段时会被误判为停用、流水线静默不跑。"""
    out = dict(acc)
    for k, v in DEFAULT_ACCOUNT_FIELDS.items():
        if k not in out or out[k] in (None, ''):
            out[k] = v if not isinstance(v, list) else list(v)
    out['active'] = bool(out.get('active', True))
    if not isinstance(out.get('tags'), list):
        out['tags'] = [t for t in str(out.get('tags') or '').split(',') if t.strip()]
    out['nickname'] = (out.get('nickname') or '').strip()
    return out

def account_artifacts(acc: dict, ledger_sheets: set[str] | None=None) -> dict:
    """单个账号的四处落点 + 数据概览。纯读，不写盘。

    ledger_sheets: 台账中存在的 sheet 名集合（由调用方统一读一次，避免逐账号 IO）。
    """
    a = _account_norm(acc)
    name = a.get('account', '')
    out_dir = Path(a['out_dir']).resolve() if a.get('out_dir') else OUTPUT_DIR
    audio_dir = Path(a['audio_dir']).resolve() if a.get('audio_dir') else OUTPUT_DIR / f'{name}_audio'
    scrape_path = out_dir / f'scrape_{name}.json'
    if not scrape_path.exists():
        alt = OUTPUT_DIR / f'scrape_{name}.json'
        scrape_path = alt if alt.exists() else scrape_path
    items, scrape_errors, last_video, scraped_at = ([], 0, None, None)
    if scrape_path.exists():
        d = read_json(scrape_path, {}) or {}
        items = d.get('items') or []
        scrape_errors = len(d.get('errors') or [])
        scraped_at = d.get('scraped_at')
        times = [i.get('publish_time_gmt8') for i in items if i.get('publish_time_gmt8')]
        last_video = max(times) if times else None
    audio_n = trans_n = 0
    if audio_dir.exists():
        audio_n = sum((1 for _ in audio_dir.glob('*.mp3')))
        tdir = audio_dir / 'transcripts'
        if tdir.exists():
            trans_n = sum((1 for _ in tdir.glob('*_transcript.txt')))
    has_sheet = bool(ledger_sheets and f'{SHEET_PREFIX}{name}' in ledger_sheets)
    if a.get('media_type') == 'image':
        has_sheet = False
        raw = load_accounts_raw()
        img = Path(raw['img_xlsx']) if raw.get('img_xlsx') else None
        if img is not None and img.is_file():
            import openpyxl
            try:
                wb = openpyxl.load_workbook(img, read_only=True)
                try:
                    if '图文账号记录' in wb.sheetnames:
                        has_sheet = any((str(row[0] or '').lstrip('@').strip() == name for row in wb['图文账号记录'].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)))
                finally:
                    wb.close()
            except Exception as e:
                raise ValueError(f'无法检查图文台账 {img}，为保护历史数据已阻止删除：{e}') from e
    age_days = None
    if last_video:
        try:
            from datetime import datetime as _dt
            lv = _dt.strptime(last_video[:19], '%Y-%m-%d %H:%M:%S')
            age_days = (_dt.now() - lv).days
        except Exception:
            age_days = None
    has_scrape = scrape_path.exists()
    has_audio = audio_n > 0
    has_any = has_scrape or has_audio or has_sheet
    if not a['active']:
        health = 'off'
    elif not has_any:
        health = 'config_only'
    elif not (has_scrape and has_sheet):
        health = 'partial'
    elif age_days is not None and age_days > STALE_DAYS:
        health = 'stale'
    else:
        health = 'ok'
    label, tone = HEALTH_LABELS[health]
    suggested = []
    if not has_scrape:
        suggested.append('s2')
    elif not has_audio and a.get('media_type') != 'image':
        suggested.append('s1')
    if a.get('media_type') != 'image' and has_audio and (trans_n < audio_n):
        suggested.append('transcribe')
    if not suggested:
        suggested = ['all']
    return {'account': name, 'nickname': a['nickname'], 'lang': a.get('lang', ''), 'active': a['active'], 'tags': a['tags'], 'note': a['note'], 'out_dir': str(out_dir), 'audio_dir': str(audio_dir), 'has_config': True, 'has_scrape': has_scrape, 'has_audio': has_audio, 'has_sheet': has_sheet, 'scrape_path': str(scrape_path) if has_scrape else None, 'video_count': len(items), 'scrape_errors': scrape_errors, 'audio_count': audio_n, 'transcript_count': trans_n, 'last_video_at': last_video, 'scraped_at': scraped_at, 'age_days': age_days, 'health': health, 'health_label': label, 'health_tone': tone, 'suggested_actions': suggested, 'missing': [k for k, ok in (('scrape', has_scrape), ('audio', has_audio or a.get('media_type') == 'image'), ('sheet', has_sheet)) if not ok]}

def accounts_with_artifacts(ledger_sheets: set[str] | None=None) -> list[dict]:
    """全部账号（含停用）+ 落点状态。UI 用，故不过滤 active。"""
    cfg = load_accounts(include_inactive=True)
    return [account_artifacts(a, ledger_sheets) for a in cfg.get('accounts', [])]

def orphan_sheets(ledger_sheets: set[str] | None=None) -> list[str]:
    """台账里存在「对标记录-xxx」但 accounts.json 无对应配置的账号名。"""
    if not ledger_sheets:
        return []
    known = {a.get('account') for a in load_accounts_raw().get('accounts', [])}
    out = []
    for s in ledger_sheets:
        if not s.startswith(SHEET_PREFIX):
            continue
        nm = s[len(SHEET_PREFIX):]
        if nm and nm not in known:
            out.append(nm)
    return sorted(out)

def orphan_scrapes() -> list[str]:
    """output 里有 scrape_xxx.json 但 accounts.json 无对应配置的账号名。"""
    known = {a.get('account') for a in load_accounts_raw().get('accounts', [])}
    if not OUTPUT_DIR.exists():
        return []
    out = []
    for p in OUTPUT_DIR.glob('scrape_*.json'):
        nm = p.stem[len('scrape_'):]
        if nm and nm not in known:
            out.append(nm)
    return sorted(out)
EDITABLE_FIELDS = ('nickname', 'lang', 'tags', 'note', 'out_dir', 'audio_dir', 'region')

def add_account(name: str, fields: dict | None=None) -> dict:
    """新增账号。name 为账号名，fields 可含 region/lang/nickname/tags/note/out_dir/audio_dir。

    重名则抛 ValueError（要改主键请走「归档旧 + 新建」）。
    ★ 2026-10-16：region→lang 自动推断 —— 若用户未显式传 lang，则按 region 推断
    （做哪个地区就用该地区语言转写）；region 未标注且未传 lang 时默认 en。
    """
    name = validate_account_name(name)
    cfg = load_accounts_raw()
    accs = cfg.setdefault('accounts', [])
    if any((isinstance(a, dict) and a.get('account') == name for a in accs)):
        raise ValueError(f'账号已存在：{name}')
    fields = validate_account_fields(fields or {})
    lang = guess_lang(fields.get('region'), fields.get('lang'))
    row = default_account(name, lang)
    for k in EDITABLE_FIELDS:
        v = fields.get(k)
        if v in (None, ''):
            continue
        if k == 'tags' and isinstance(v, str):
            v = [t.strip() for t in v.split(',') if t.strip()]
        row[k] = v
    row['active'] = bool(fields.get('active', True))
    accs.append(row)
    _write_accounts(cfg)
    return row
MEDIA_TYPES = ('video', 'image')
IMAGE_TAG = '图文账号'

def set_media_type(accounts, media_type: str='image') -> dict:
    """批量把账号设为图文账号 / 视频账号（一次性原子写）。

    image：写 media_type="image" 并补「图文账号」标签 —— 流水线据此跳过音频/转写，
           走图文枚举（enumerate_image.py）。
    video：移除 media_type 与「图文账号」标签，回到视频口径。
    只改账号类型，不动台账 sheet 与任何产物；逐条返回 saved/same/not_found/invalid。
    """
    mt = str(media_type or '').strip().lower()
    if mt not in MEDIA_TYPES:
        raise ValueError(f"media_type 只能是 {' / '.join(MEDIA_TYPES)}")
    if isinstance(accounts, (str, bytes)) or not isinstance(accounts, (list, tuple)):
        raise ValueError('accounts 必须是数组')
    cfg = load_accounts_raw()
    index = {a.get('account'): a for a in cfg.get('accounts', []) if isinstance(a, dict)}
    results, saved = ([], [])
    for raw in accounts:
        name = raw.strip() if isinstance(raw, str) else ''
        if not name:
            results.append({'account': '', 'status': 'invalid', 'reason': '账号名为空'})
            continue
        row = index.get(name)
        if row is None:
            results.append({'account': name, 'status': 'not_found', 'reason': 'accounts.json 中不存在'})
            continue
        cur = str(row.get('media_type') or 'video').strip().lower()
        tags = row.get('tags') or []
        if not isinstance(tags, list):
            tags = [t.strip() for t in str(tags).split(',') if t.strip()]
        if mt == 'image':
            if cur == 'image' and IMAGE_TAG in tags:
                results.append({'account': name, 'status': 'same', 'reason': '已经是图文账号'})
                continue
            row['media_type'] = 'image'
            if IMAGE_TAG not in tags:
                tags.append(IMAGE_TAG)
            results.append({'account': name, 'status': 'saved', 'from': cur})
        else:
            if cur != 'image' and IMAGE_TAG not in tags:
                results.append({'account': name, 'status': 'same', 'reason': '已经是视频账号'})
                continue
            row.pop('media_type', None)
            tags = [t for t in tags if t != IMAGE_TAG]
            results.append({'account': name, 'status': 'saved', 'from': 'image'})
        row['tags'] = tags
        saved.append(name)
    if saved:
        _write_accounts(cfg)
    return {'media_type': mt, 'saved': saved, 'results': results}

def update_account(name: str, fields: dict) -> dict:
    """改账号的可编辑字段。主键 account 只读，试图修改直接忽略。"""
    validate_account_name(name)
    fields = validate_account_fields(fields)
    cfg = load_accounts_raw()
    for a in cfg.get('accounts', []):
        if not isinstance(a, dict) or a.get('account') != name:
            continue
        for k in EDITABLE_FIELDS:
            if k not in fields:
                continue
            v = fields[k]
            if k == 'tags' and isinstance(v, str):
                v = [t.strip() for t in v.split(',') if t.strip()]
            a[k] = v
        if 'region' in fields and fields.get('lang') in (None, ''):
            a['lang'] = guess_lang(a.get('region'))
        elif 'lang' in fields and fields['lang'] in (None, ''):
            a['lang'] = guess_lang(a.get('region'))
        if 'active' in fields:
            a['active'] = bool(fields['active'])
        _write_accounts(cfg)
        return a
    raise ValueError(f'账号不存在：{name}')

def set_account_active(name: str, active: bool) -> dict:
    return update_account(name, {'active': bool(active)})

def remove_account(name: str) -> None:
    """删除账号（仅限无任何产物的空壳，避免丢失历史索引）。"""
    validate_account_name(name)
    cfg = load_accounts_raw()
    accs = cfg.get('accounts', [])
    row = next((a for a in accs if isinstance(a, dict) and a.get('account') == name), None)
    if row is None:
        raise ValueError(f'账号不存在：{name}')
    ledger_sheets: set[str] = set()
    ledger = ledger_path()
    if ledger:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(ledger, read_only=True)
            ledger_sheets = set(wb.sheetnames)
            wb.close()
        except Exception as e:
            raise ValueError(f'无法检查台账 {ledger}，为保护历史数据已阻止删除：{e}') from e
    art = account_artifacts(row, ledger_sheets)
    if art['has_scrape'] or art['has_audio'] or art['has_sheet']:
        have = [k for k, ok in (('抓取', art['has_scrape']), ('音频', art['has_audio']), ('台账', art['has_sheet'])) if ok]
        raise ValueError(f"{name} 已有产物（{'/'.join(have)}），请改用「停用」保留历史索引")
    cfg['accounts'] = [a for a in accs if a is not row]
    _write_accounts(cfg)

def load_automations() -> list[dict]:
    return []

def _is_archive_ledger(path: Path) -> bool:
    parts = path.stem.lower().replace('-', '_').replace(' ', '_').split('_')
    return any((part in {'归档', 'archive', 'archived'} for part in parts))

def list_ledgers() -> list[Path]:
    """活动 xlsx（排除 ~$ 临时文件和归档簿），按修改时间倒序。"""
    if not LEDGER_DIR.exists():
        return []
    files = [p for p in LEDGER_DIR.glob('*.xlsx') if p.is_file() and (not p.name.startswith('~$')) and (not _is_archive_ledger(p))]
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)

def ledger_path() -> Path | None:
    """视频主台账：优先 accounts.json 的 xlsx 字段；缺失时退回目录内最近修改的 xlsx。
    图文台账走独立文件（accounts.json.img_xlsx），不在此返回——避免图文文件 mtime 更新后
    顶掉视频主簿、破坏 _ledger_id_sets 的视频入账覆盖判定。"""
    files = list_ledgers()
    cfg = read_json(ACCOUNTS_JSON, {}) or {}
    img = Path(cfg['img_xlsx']).resolve() if cfg.get('img_xlsx') else None
    files = [f for f in files if f.resolve() != img and (not f.stem.startswith('图文'))]
    configured = Path(cfg['xlsx']).resolve() if cfg.get('xlsx') else None
    selected = next((f for f in files if f.resolve() == configured), None)
    if selected is not None:
        return selected
    return files[0] if files else None

def image_ledger_path() -> Path | None:
    cfg = read_json(ACCOUNTS_JSON, {}) or {}
    configured = Path(cfg['img_xlsx']).resolve() if cfg.get('img_xlsx') else None
    return next((f for f in list_ledgers() if f.resolve() == configured), None)

def run_env() -> dict:
    """子进程环境：注入代理，禁用 HF 联网检查。"""
    env = dict(os.environ)
    env['HTTPS_PROXY'] = PROXY
    env['HTTP_PROXY'] = PROXY
    env['HF_HUB_OFFLINE'] = '1'
    env['PYTHONIOENCODING'] = 'utf-8'
    return env

def diagnostics() -> dict:
    """环境自检：路径与关键文件是否存在。"""
    return {'pipeline_dir': str(PIPELINE_DIR), 'pipeline_exists': PIPELINE_DIR.exists(), 'python': python_exe(), 'python_exists': VENV_PYTHON.exists(), 'accounts_json': str(ACCOUNTS_JSON), 'accounts_exists': ACCOUNTS_JSON.exists(), 'output_exists': OUTPUT_DIR.exists(), 'reports_exists': REPORTS_DIR.exists(), 'cookie_exists': COOKIE_FILE.exists(), 'ledger_dir': str(LEDGER_DIR), 'ledger': str(ledger_path()) if ledger_path() else None, 'ledger_files': [p.name for p in list_ledgers()], 'proxy': PROXY}
