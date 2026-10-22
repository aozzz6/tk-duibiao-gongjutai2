"""对标分析聚合层（第二版 · A 方案）。

把散在本地、彼此分离的三类数据归一成一套可供「对标分析」视图直接渲染的结构：
  - 视频抓取   pipeline/output/scrape_<account>.json   （stats / desc / publish_time_gmt8）
  - 小黄车商品 D:/1/对标账号/<account>/<account>_products.json  （视频 id → 商品）
  - 台账       D:/1/对标账号/*.xlsx                          （行对账 / 出单列现状）

设计原则（对应豆包「自动归一：汇总、去重、关联」）：
  - 关联：以 视频 id 为主键，把 scrape 视频 与 products 商品表 join；
  - 去重：同账号同一视频 id 只算一次；多账号不重复计总；
  - 汇总：账号维度、全局维度分别聚合。

诚实边界（对应「可借鉴点」）：
  - 出单量 / GMV：FastMoss 积分已耗尽（MCP 实测被拒），一律标注为「待接入」，绝不编造数字。
  - 「爆款规律 / 可借鉴点」由规则从真实播放/互动/标签/钩子/商品覆盖推导，
    属启发式，前端显式标注。
"""
from __future__ import annotations
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import copy
import glob
import json
import os
import re
import time
from collections import Counter
from pathlib import Path
import config
_cache: dict = {'ts': 0.0, 'data': None}
_TAG_RE = re.compile('#(\\w[\\w-]+)')
_FASTMOSS_PRODUCT_URL = 'https://www.fastmoss.com/e-commerce/detail/{}'
_PRODUCTS_TOP_N = 10
_PRODUCT_VIDEOS_N = 5

def _num(v) -> int:
    """stats 里 collectCount 可能是字符串，统一转 int。"""
    if v is None:
        return 0
    try:
        return int(float(v))
    except (ValueError, TypeError, OverflowError):
        return 0

def _load_scrape(account: str) -> list[dict]:
    p = config.OUTPUT_DIR / f'scrape_{account}.json'
    if not p.exists():
        return []
    data = config.read_json(p, None) or {}
    return data.get('items', []) or []

def _load_products_map(account: str) -> dict:
    """视频 id → [{product_id,title,...}]。"""
    p = next((p for p in (config.LEDGER_DIR / account / f'{account}_products.json', config.OUTPUT_DIR / f'{account}_products.json') if p.exists()), None)
    if p is None:
        return {}
    data = config.read_json(p, None) or {}
    if isinstance(data, dict) and isinstance(data.get('products'), dict):
        data = data['products']
    if isinstance(data, dict):
        return {str(k): [p for p in (v if isinstance(v, list) else [v]) if isinstance(p, dict) and p] for k, v in data.items()}
    return {}
_LEDGER_CACHE: dict = {'key': None, 'data': {}}
_LEDGER_FIELDS = {'time': ('发布时',), 'title': ('标题',), 'tags': ('话题标签',), 'link': ('视频链接',), 'plays': ('播放量',), 'likes': ('点赞',), 'comments': ('评论',), 'shares': ('转发',), 'collects': ('收藏',), 'audio': ('音频',), 'trans': ('翻译',), 'acct': ('对标账号',), 'product': ('商品链接',)}

def _col_index(header: list[str], names: tuple, default: int | None=None) -> int | None:
    """按表头名定位列（不写死序号，列序调整不会错位）。"""
    for i, h in enumerate(header):
        if h and any((n in h for n in names)):
            return i
    return default

def _ledger_files() -> list[tuple[Path, str]]:
    """(路径, 媒体类型)：视频主簿 + 图文独立簿（两本口径不同，12/14 列）。"""
    out = []
    for p, mt in ((config.ledger_path(), 'video'), (config.image_ledger_path(), 'image')):
        if p and Path(p).exists():
            out.append((Path(p), mt))
    return out

def _ledger_rows() -> dict[str, list[dict]]:
    """读台账 → {account: [row]}。缓存键 = 文件 mtime+size，避免每 15s 重开 xlsx。"""
    files = _ledger_files()
    key = tuple(((str(p), p.stat().st_mtime, p.stat().st_size) for p, _ in files))
    if _LEDGER_CACHE['key'] == key:
        return _LEDGER_CACHE['data']
    import openpyxl
    data: dict[str, list[dict]] = {}
    seen_by_acc: dict[str, set] = {}
    for path, media_type in files:
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            for sn in wb.sheetnames:
                if sn.startswith('说明') or sn.startswith('对标账号记录-'):
                    continue
                ws = wb[sn]
                rows = ws.iter_rows(values_only=True)
                header = [str(c or '').strip() for c in next(rows, ())]
                if not header:
                    continue
                ix = {k: _col_index(header, v) for k, v in _LEDGER_FIELDS.items()}
                if ix['link'] is None:
                    continue
                sheet_acc = sn.split('-', 1)[1] if sn.startswith('对标记录-') else ''
                for row in rows:
                    cell = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
                    link = str(cell('link') or '')
                    m = re.search('/video/(\\d{15,25})', link) or re.search('ID[:：]\\s*(\\d{15,25})', link) or re.search('(\\d{15,25})', link)
                    if not m:
                        continue
                    vid = m.group(1)
                    acc = sheet_acc or str(cell('acct') or '').strip().lstrip('@')
                    if not acc:
                        continue
                    seen = seen_by_acc.setdefault(acc, set())
                    if vid in seen:
                        continue
                    seen.add(vid)
                    plays = _num(cell('plays'))
                    likes = _num(cell('likes'))
                    cmt = _num(cell('comments'))
                    shr = _num(cell('shares'))
                    col = _num(cell('collects'))
                    title = str(cell('title') or '').strip()
                    tags = str(cell('tags') or '').strip()
                    data.setdefault(acc, []).append({'vid': vid, 'title': _first_line(title), 'desc': (title + (' ' + tags if tags else '')).strip(), 'plays': plays, 'likes': likes, 'comments': cmt, 'shares': shr, 'collects': col, 'publish_time': str(cell('time') or '').strip(), 'url': link if link.startswith('http') else f'https://www.tiktok.com/@{acc}/video/{vid}', 'products': [], 'product_link': str(cell('product') or '').strip(), 'engagement': round((likes + cmt + shr + col) / plays * 100, 3) if plays else 0.0, 'source': 'ledger', 'media_type': media_type, 'has_audio': bool(str(cell('audio') or '').strip()), 'has_translation': bool(str(cell('trans') or '').strip())})
        finally:
            wb.close()
    _LEDGER_CACHE.update(key=key, data=data)
    return data

def _first_line(desc: str, limit: int=120) -> str:
    s = (desc or '').strip()
    s = re.sub('\\s+', ' ', s)
    return s[:limit]

def _hooks(desc: str, n: int=4) -> list[str]:
    """文案开头词（去掉话题标签），作为「钩子」近似。"""
    s = (desc or '').strip()
    s = _TAG_RE.sub(' ', s)
    words = re.findall("[A-Za-z0-9']+", s)
    stop = {'the', 'a', 'an', 'my', 'our', 'in', 'for', 'with', 'to', 'and', 'i', 'is', 'are', 'on', 'this', 'that'}
    out = []
    for w in words:
        wl = w.lower()
        if wl in stop or len(wl) < 3:
            continue
        out.append(w)
        if len(out) >= n:
            break
    return out

def _artifact_present(acc: str, accounts_cfg: list[dict]) -> bool:
    """账号是否有任一产物（与总览 _account_view 的 active 判定同口径）：
    音频 / 转写 / 抓取 JSON / 商品文件，任一存在即 True。"""
    row = next((a for a in accounts_cfg if a.get('account') == acc), None)
    try:
        audio_dir = Path(row['audio_dir']) if row and row.get('audio_dir') else config.OUTPUT_DIR / f'{acc}_audio'
        out_dir = Path(row['out_dir']) if row and row.get('out_dir') else config.OUTPUT_DIR
    except (TypeError, ValueError):
        audio_dir, out_dir = (config.OUTPUT_DIR / f'{acc}_audio', config.OUTPUT_DIR)
    if any(audio_dir.glob(f'{acc}_*.mp3')) or any(audio_dir.glob(f'{acc}_*.m4a')):
        return True
    if any((audio_dir / 'transcripts').glob(f'{acc}_*_transcript.txt')):
        return True
    if (out_dir / f'scrape_{acc}.json').exists():
        return True
    for cand in (config.LEDGER_DIR / acc / f'{acc}_products.json', out_dir / f'{acc}_products.json', audio_dir / f'{acc}_products.json'):
        if cand.exists():
            return True
    return False

def _aggregate():
    accounts_cfg = config.load_accounts(include_inactive=True).get('accounts', [])
    acc_names = [a.get('account') for a in accounts_cfg if a.get('account')]
    for f in glob.glob(str(config.LEDGER_DIR / '*' / '*_products.json')):
        acc = Path(f).parent.name
        if acc not in acc_names:
            acc_names.append(acc)
    video_map: dict[str, list[dict]] = {}
    products_map: dict[str, dict] = {}
    for acc in acc_names:
        items = _load_scrape(acc)
        prods = _load_products_map(acc)
        vids = []
        for it in items:
            vid = str(it.get('id') or it.get('video_id') or '')
            stats = it.get('stats', {}) or {}
            rec = {'vid': vid, 'title': _first_line(it.get('desc', '')), 'desc': (it.get('desc') or '').strip(), 'plays': _num(stats.get('playCount')), 'likes': _num(stats.get('diggCount')), 'comments': _num(stats.get('commentCount')), 'shares': _num(stats.get('shareCount')), 'collects': _num(stats.get('collectCount')), 'publish_time': it.get('publish_time_gmt8', ''), 'url': it.get('url', ''), 'products': prods.get(vid, []) if vid else [], 'source': 'scrape', 'media_type': 'video'}
            denom = rec['plays']
            rec['engagement'] = round((rec['likes'] + rec['comments'] + rec['shares'] + rec['collects']) / denom * 100, 3) if denom else 0.0
            vids.append(rec)
        seen = set()
        uniq = []
        for v in vids:
            if v['vid'] and v['vid'] in seen:
                continue
            if v['vid']:
                seen.add(v['vid'])
            uniq.append(v)
        video_map[acc] = uniq
        products_map[acc] = prods
    account_rows = []
    for acc in acc_names:
        vids = video_map.get(acc, [])
        if not vids:
            continue
        total_plays = sum((v['plays'] for v in vids))
        total_likes = sum((v['likes'] for v in vids))
        with_prod = sum((1 for v in vids if v['products']))
        tags = Counter()
        for v in vids:
            for t in _TAG_RE.findall(v['desc'] or ''):
                tags[t.lower()] += 1
        top_tags = [t for t, _ in tags.most_common(5)]
        hooks = Counter()
        for v in vids:
            for h in _hooks(v['desc'], 3):
                hooks[h.lower()] += 1
        top_hooks = [h for h, _ in hooks.most_common(3)]
        avg_eng = sum((v['engagement'] for v in vids)) / len(vids) if vids else 0
        account_rows.append({'account': acc, 'lang': next((a.get('lang') for a in accounts_cfg if a.get('account') == acc), ''), 'videos': len(vids), 'total_plays': total_plays, 'total_likes': total_likes, 'avg_plays': round(total_plays / len(vids)) if vids else 0, 'avg_engagement': round(avg_eng, 2), 'with_products': with_prod, 'product_coverage': round(100 * with_prod / len(vids), 1) if vids else 0, 'top_tags': top_tags, 'top_hooks': top_hooks, 'has_scrape': True})
    account_rows.sort(key=lambda r: r['total_plays'], reverse=True)
    all_vids = [v for acc in acc_names for v in video_map.get(acc, [])]
    for acc in acc_names:
        for v in video_map.get(acc, []):
            v['account'] = acc
    kpis, patterns = _derive(all_vids, account_rows, products_of=lambda v: v.get('products') or [])
    takeaways = _build_takeaways(kpis, patterns, account_rows)
    ledger_rows = _ledger_rows()
    ledger_only: list[dict] = []
    for acc, rows in ledger_rows.items():
        have = {v['vid'] for v in video_map.get(acc, [])}
        for r in rows:
            if r['vid'] in have:
                continue
            have.add(r['vid'])
            ledger_only.append({**r, 'account': acc})
    opt_names = [acc for acc in acc_names if video_map.get(acc) or _artifact_present(acc, accounts_cfg)]
    for acc in ledger_rows:
        if acc not in opt_names:
            opt_names.append(acc)
    account_options = [{'account': acc, 'videos': len(video_map.get(acc) or [])} for acc in opt_names]
    kpis['ledger_rows'] = sum((len(v) for v in ledger_rows.values()))
    kpis['ledger_only'] = len(ledger_only)
    videos_out = _flatten_videos(video_map) + ledger_only
    videos_out.sort(key=lambda v: v.get('plays') or 0, reverse=True)
    return {'ts': time.time(), 'kpis': kpis, 'accounts': account_rows, 'account_options': account_options, 'videos': videos_out, 'patterns': patterns, 'takeaways': takeaways, 'data_sources': {'scrape_dir': str(config.OUTPUT_DIR), 'products_dir': str(config.LEDGER_DIR), 'ledger_rows': sum((len(v) for v in ledger_rows.values())), 'ledger_only': len(ledger_only), 'gmv_status': 'pending（FastMoss 积分已耗尽，视频级出单/GMV 待补查）'}}

def _flatten_videos(video_map) -> list[dict]:
    out = []
    for acc, vids in video_map.items():
        for v in vids:
            row = dict(v)
            row['account'] = acc
            out.append(row)
    out.sort(key=lambda v: v['plays'], reverse=True)
    return out

def _build_takeaways(kpis, patterns, account_rows) -> list[dict]:
    """由真实数据推导的可借鉴点。每条 = {text, basis, confidence}。"""
    out = []
    if not kpis['videos']:
        return out
    buckets = sorted(patterns['play_buckets'], key=lambda b: b['pct'], reverse=True)
    if buckets:
        top = buckets[0]
        out.append({'text': f"{top['pct']}% 的视频落在 {top['label']} 档（共 {top['count']} 条）——爆款呈长尾，多数视频跑不出量，少数头部贡献主要曝光。", 'basis': '播放量分桶统计', 'confidence': '数据'})
    cov = kpis['product_coverage']
    if cov >= 90:
        out.append({'text': f"挂车率 {cov}%（{kpis['with_products']}/{kpis['videos']}），几乎每条都带小黄车——对标账号以「内容挂车直接转化」为主打法。", 'basis': '商品覆盖统计', 'confidence': '数据'})
    elif cov and cov < 90:
        out.append({'text': f'挂车率 {cov}%，部分视频纯内容不带车。高互动但未挂车的视频是可补挂车的选品机会。', 'basis': '商品覆盖统计', 'confidence': '数据'})
    if patterns['top_tags']:
        tags = ', '.join(('#' + t['tag'] for t in patterns['top_tags'][:6]))
        out.append({'text': f'高频话题标签：{tags}。这些是对标赛道当前投流 / 起量的关键词，可直接复用进自有账号的标签矩阵。', 'basis': '标签词频', 'confidence': '数据'})
    if patterns['hook_words']:
        hooks = ', '.join((w for w in patterns['hook_words'][:6]))
        out.append({'text': f'开场高频词：{hooks}。视频前 3 秒普遍用悬念/功能/痛点词勾住观众。', 'basis': '文案开头词频', 'confidence': '启发式'})
    if patterns['products_top']:
        p = patterns['products_top'][0]
        nm = (p.get('title') or p.get('product_id') or '未知商品').strip()
        out.append({'text': f"被挂车最多的商品：「{nm[:40]}…」（出现在 {p['vids']} 条视频）。是账号主推款，可重点拆解其卖点话术。", 'basis': '商品出现频次', 'confidence': '数据'})
    if account_rows:
        top_acc = account_rows[0]
        out.append({'text': f"播放量最高账号 {top_acc['account']}：{top_acc['videos']} 条视频、总播放 {top_acc['total_plays']:,}、平均互动率 {top_acc['avg_engagement']}%。其钩子「{', '.join(top_acc['top_hooks'][:3]) or '-'}」可优先模仿。", 'basis': '账号维度聚合', 'confidence': '数据'})
    return out

def _derive(all_vids: list[dict], account_rows: list[dict], products_of=None) -> tuple[dict, dict]:
    """由一组视频推导 KPI 与 patterns（供全局聚合与账号过滤共用，保证口径一致）。

    products_of(vid) -> list[product] 用于按账号子集重算「被挂车商品」；
    不传则退化为直接读 video["products"]。
    """
    if products_of is None:
        products_of = lambda v: v.get('products') or []
    total_plays = sum((v['plays'] for v in all_vids))
    with_prod = sum((1 for v in all_vids if products_of(v)))
    avg_eng = sum((v['engagement'] for v in all_vids)) / len(all_vids) if all_vids else 0
    kpis = {'accounts': len([r for r in account_rows if r['videos']]), 'videos': len(all_vids), 'total_plays': total_plays, 'total_likes': sum((v['likes'] for v in all_vids)), 'avg_engagement': round(avg_eng, 2), 'with_products': with_prod, 'product_coverage': round(100 * with_prod / len(all_vids), 1) if all_vids else 0, 'top_accounts_by_plays': [r['account'] for r in account_rows[:3]]}
    buckets = [('0-1k', 0, 1000), ('1k-10k', 1000, 10000), ('10k-100k', 10000, 100000), ('100k+', 100000, float('inf'))]
    bucket_counts = {label: 0 for label, _, _ in buckets}
    for v in all_vids:
        for label, lo, hi in buckets:
            if lo <= v['plays'] < hi:
                bucket_counts[label] += 1
                break
    play_buckets = [{'label': label, 'count': bucket_counts[label], 'pct': round(100 * bucket_counts[label] / len(all_vids), 1) if all_vids else 0} for label, _, _ in buckets]
    gtags = Counter()
    for v in all_vids:
        for t in _TAG_RE.findall(v['desc'] or ''):
            gtags[t.lower()] += 1
    top_tags = [{'tag': t, 'count': c, 'pct': round(100 * c / len(all_vids), 1) if all_vids else 0} for t, c in gtags.most_common(12)]
    ghooks = Counter()
    for v in all_vids:
        for h in _hooks(v['desc'], 3):
            ghooks[h.lower()] += 1
    hook_words = [w for w, _ in ghooks.most_common(12)]
    prod_map: dict[str, dict] = {}
    for v in all_vids:
        for p in products_of(v):
            if not isinstance(p, dict):
                continue
            t = (p.get('title') or '').strip()
            pid = str(p.get('product_id') or '').strip()
            if not t and (not pid):
                continue
            key = pid or 'title:' + t[:80]
            rec = prod_map.get(key)
            if rec is None:
                rec = {'product_id': pid, 'title': t, 'anchor_type': p.get('anchor_type'), 'fastmoss_url': _FASTMOSS_PRODUCT_URL.format(pid) if pid else '', 'vids': 0, 'accounts': [], 'total_plays': 0, '_videos': []}
                prod_map[key] = rec
            rec['vids'] += 1
            if not rec['title'] and t:
                rec['title'] = t
            acc = v.get('account') or ''
            if acc and acc not in rec['accounts']:
                rec['accounts'].append(acc)
            rec['total_plays'] += int(v.get('plays') or 0)
            rec['_videos'].append({'account': acc, 'vid': v.get('vid', ''), 'title': v.get('title') or '', 'plays': int(v.get('plays') or 0), 'url': v.get('url') or '', 'publish_time': v.get('publish_time') or ''})
    ranked = sorted(prod_map.values(), key=lambda r: (-r['vids'], -r['total_plays']))
    products_top = []
    for r in ranked[:_PRODUCTS_TOP_N]:
        r['videos'] = sorted(r.pop('_videos'), key=lambda x: x['plays'], reverse=True)[:_PRODUCT_VIDEOS_N]
        r['pct'] = round(100 * r['vids'] / with_prod, 1) if with_prod else 0
        products_top.append(r)
    top5 = sorted(all_vids, key=lambda v: v['plays'], reverse=True)[:5]
    top_videos = [{'account': v.get('account', ''), 'vid': v['vid'], 'title': v['title'], 'plays': v['plays'], 'url': v['url']} for v in top5]
    patterns = {'play_buckets': play_buckets, 'top_tags': top_tags, 'hook_words': hook_words, 'top_videos': top_videos, 'products_top': products_top, 'n_videos_in_sample': len(all_vids)}
    return (kpis, patterns)

def _deps_mtime() -> float:
    """benchmark 依赖文件的最新 mtime：accounts.json + 抓取产物 + 商品文件。
    任一变化都应使聚合缓存失效（账号增删 / 新抓取 / 商品提取后立即生效，
    与总览看板保持同步，而不是等 15s TTL 自然过期）。"""
    mt = 0.0
    if config.ACCOUNTS_JSON.exists():
        mt = max(mt, config.ACCOUNTS_JSON.stat().st_mtime)
    for pat in ('scrape_*.json',):
        for f in glob.glob(str(config.OUTPUT_DIR / pat)):
            try:
                mt = max(mt, os.path.getmtime(f))
            except OSError:
                pass
    for f in glob.glob(str(config.LEDGER_DIR / '*' / '*_products.json')):
        try:
            mt = max(mt, os.path.getmtime(f))
        except OSError:
            pass
    for p, _ in _ledger_files():
        try:
            mt = max(mt, p.stat().st_mtime)
        except OSError:
            pass
    return mt

def build(account: str | None=None, refresh: bool=False, accounts: list[str] | None=None) -> dict:
    deps = _deps_mtime()
    stale = deps > _cache.get('deps', 0.0)
    if not refresh and (not stale) and _cache['data'] and (time.time() - _cache['ts'] < 15):
        data = _cache['data']
    else:
        data = _aggregate()
        _cache.update(ts=time.time(), data=data, deps=max(deps, time.time()))
    data = copy.deepcopy(data)
    if accounts is not None or account:
        selected = list(dict.fromkeys((a.strip() for a in accounts if a.strip()))) if accounts is not None else [account]
        keep = set(selected)
        data['videos'] = [v for v in data['videos'] if v['account'] in keep]
        data['accounts'] = [a for a in data['accounts'] if a['account'] in keep]
        scrape_only = [v for v in data['videos'] if v.get('source') != 'ledger']
        kpis, patterns = _derive(scrape_only, data['accounts'], products_of=lambda v: v.get('products') or [])
        data['kpis'] = kpis
        data['patterns'] = patterns
        data['takeaways'] = _build_takeaways(kpis, patterns, data['accounts'])
        data['filtered_account'] = account
        if accounts is not None:
            data['filtered_accounts'] = selected
    return data
