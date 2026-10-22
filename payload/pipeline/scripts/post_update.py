import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import argparse
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path
PIPE = Path(__file__).resolve().parent.parent
PIPE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from desc_split import COL_NOTE, COL_NOTE_ZH
from build_ledger_15col import valid_header, _valid_img_header, vid_of
PY = sys.executable
BUILDER = PIPE / 'scripts' / 'build_ledger_15col.py'
CFG_PATH = PIPE / 'accounts.json'
WORKBENCH_URL = os.environ['WB_WORKBENCH_URL'] + '/api/overview?refresh=true'

def _cell(v) -> str:
    """单元格取文本（None → 空串，非字符串先转字符串再 strip）。"""
    return '' if v is None else str(v).strip()

def log(msg: str) -> None:
    print(f'[{dt.datetime.now():%H:%M:%S}] {msg}', flush=True)

def load_cfg() -> dict:
    return json.loads(CFG_PATH.read_text(encoding='utf-8'))

def run_builder(args: list[str]) -> int:
    """委托 build_ledger_15col.py（单一事实源，本脚本不复制其逻辑）。"""
    cmd = [PY, str(BUILDER)] + args
    log('运行: ' + ' '.join(cmd[2:]))
    proc = subprocess.run(cmd)
    return proc.returncode

def check_api(scope: list[str] | None=None, cfg: dict | None=None) -> dict:
    """工作台 overview 校验。服务不在线返回 skipped（不算失败）。

    ★ 2026-10-21 验收 B（停用账号顶红）：旧实现 scope=None 时遍历 overview 全部账号，
      与 check_ledger 的 active/ledger_skip 口径不一致 —— 一个停用账号
      （accounts.json `active:false`，overview 里 `active:true, enabled:false`）
      就能把 check_api 判成 error → post_update rc=1 → 日更的 ledger_verify 恒 ERROR。

      注意 overview 的 `active` 是「有产物」语义，不是配置开关；这里必须按
      accounts.json 的 **active + ledger_skip** 计算「启用名单」，再限定终验范围。
      停用账号的历史产物仍在 overview/台账里可见，只是不再参与终验判定。
    """
    try:
        import urllib.request
        with urllib.request.urlopen(urllib.request.Request(WORKBENCH_URL, headers={'X-Workbench-Token': os.environ['WB_DESKTOP_TOKEN']}), timeout=15) as r:
            data = json.loads(r.read().decode('utf-8'))
    except Exception as e:
        log(f'[warn] 工作台不可达（{str(e)[:80]}），跳过 API 校验，仅磁盘核对')
        return {'status': 'skipped', 'reason': str(e)[:120]}
    accounts = data.get('accounts') or []
    cfg = cfg or load_cfg()
    enabled = {a.get('account') for a in cfg.get('accounts', []) if a.get('active', True) and (not a.get('ledger_skip', False))}
    seen = len(accounts)
    disabled = sorted({str(a.get('account')) for a in accounts if a.get('account') and a.get('account') not in enabled})
    accounts = [a for a in accounts if a.get('account') in enabled]
    if scope is not None:
        accounts = [a for a in accounts if a.get('account') in scope]
    bad = []
    for a in accounts:
        name = a.get('account', '?')
        sev = a.get('severity')
        gaps = a.get('gaps') or []
        if sev not in (None, 'ok'):
            bad.append({'account': name, 'severity': sev, 'gaps': [f"{g.get('label')}:{g.get('msg')}" for g in gaps]})
    return {'status': 'ok' if not bad else 'error', 'bad': bad, 'checked': len(accounts), 'scope_basis': 'accounts.json active+ledger_skip', 'skipped_disabled': disabled, 'overview_total': seen}

def check_ledger(cfg: dict | None=None, scope: list[str] | None=None) -> dict:
    """openpyxl 复读台账：行数对账 + N 列空缺清单。纯只读。"""
    import openpyxl
    cfg = cfg or load_cfg()
    accounts = [a for a in cfg.get('accounts', []) if a.get('active', True) and (not a.get('ledger_skip', False))]
    if scope is not None:
        known = {a['account'] for a in accounts}
        if any((a not in known for a in scope)):
            return {'status': 'error', 'error': '指定账号不在台账范围'}
        accounts = [a for a in accounts if a['account'] in scope]
    need_video = any((a.get('media_type') != 'image' for a in accounts))
    need_image = any((a.get('media_type') == 'image' for a in accounts))
    ledger = Path(cfg.get('xlsx') or '')
    if need_video and (not ledger.is_file()):
        return {'status': 'error', 'error': f'视频台账不存在: {ledger}'}
    img_ledger = Path(cfg.get('img_xlsx') or '')
    if need_image and (not cfg.get('img_xlsx') or img_ledger.resolve() == ledger.resolve()):
        return {'status': 'error', 'error': 'img_xlsx 必须为独立图文文件'}
    if need_image and (not img_ledger.is_file()):
        return {'status': 'error', 'error': f'图文台账不存在: {img_ledger}'}
    targets = ([ledger] if need_video else []) + ([img_ledger] if need_image else [])
    if any((list(p.parent.glob('~$*')) for p in targets)):
        return {'status': 'error', 'error': '台账被占用（~$ 锁文件存在，Excel/WPS 未关闭？）'}
    wb = openpyxl.load_workbook(ledger, read_only=True) if need_video else None
    img_wb = openpyxl.load_workbook(img_ledger, read_only=True) if need_image else None
    details, bad = ([], [])
    for acc_cfg in accounts:
        acc = acc_cfg['account']
        is_img = acc_cfg.get('media_type') == 'image'
        sn = '图文账号记录' if is_img else f'对标记录-{acc}'
        ws_src = img_wb if is_img and img_wb is not None else wb
        out = Path(acc_cfg.get('out_dir') or PIPE / 'output')
        sp = out / f'scrape_{acc}.json'
        if sn not in ws_src.sheetnames:
            details.append({'account': acc, 'sheet': 'MISSING', 'rows': 0, 'scrape_items': None, 'n_empty': None, 'media_type': 'image' if is_img else 'video'})
            bad.append(acc)
            continue
        if not sp.exists():
            details.append({'account': acc, 'status': 'skipped', 'reason': 'scrape artifact missing', 'rows': None, 'scrape_items': None})
            continue
        try:
            expected = {str(i['id']) for i in json.loads(sp.read_text(encoding='utf-8')).get('items', []) if i.get('id')}
            items_n = len(expected)
        except Exception as e:
            details.append({'account': acc, 'status': 'error', 'reason': f'scrape artifact unreadable: {e}', 'rows': None, 'scrape_items': None})
            bad.append(acc)
            continue
        ws = ws_src[sn]
        rows, n_empty, actual = (0, [], [])
        header_ok = _valid_img_header(ws) if is_img else valid_header(ws)
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            if is_img:
                if len(row) < 2 or str(row[1] or '').lstrip('@').strip() != acc:
                    continue
                rows += 1
                v = vid_of(row[5]) if len(row) > 5 else None
                if v:
                    actual.append(v)
                continue
            rows += 1
            v = vid_of(row[5]) if len(row) > 5 else None
            if v:
                actual.append(v)
            padded = list(row) + [None] * (COL_NOTE_ZH - len(row))
            m, n = (_cell(padded[COL_NOTE - 1]), _cell(padded[COL_NOTE_ZH - 1]))
            if m and (not n):
                n_empty.append(i)
        short = rows < items_n
        missing = sorted(expected - set(actual))
        duplicates = len(actual) - len(set(actual))
        details.append({'account': acc, 'rows': rows, 'scrape_items': items_n, 'rows_short': short, 'n_empty_count': len(n_empty), 'n_empty_rows': n_empty[:20], 'missing_ids': missing, 'duplicate_ids': duplicates, 'header_ok': header_ok, 'media_type': 'image' if is_img else 'video'})
        if short or missing or duplicates or (not header_ok):
            bad.append(acc)
    if wb is not None:
        wb.close()
    if img_wb is not None:
        img_wb.close()
    return {'status': 'ok' if not bad else 'error', 'bad': bad, 'accounts': details}

def main_verify(skip_api: bool, scope: list[str] | None=None) -> int:
    api = {'status': 'skipped', 'reason': '--skip-api'} if skip_api else check_api(scope) if scope is not None else check_api()
    ledger = check_ledger(scope=scope) if scope is not None else check_ledger()
    level = 'ERROR' if 'error' in (api.get('status'), ledger.get('status')) else 'INFO'
    print(json.dumps({'level': level, 'api': api, 'ledger': ledger, 'date': dt.date.today().isoformat()}, ensure_ascii=False, indent=1))
    if ledger.get('status') == 'ok':
        for d in ledger['accounts']:
            extra = f" N空缺={d['n_empty_count']}" if d.get('n_empty_count') is not None else ''
            warn = ' ⚠行数不足' if d.get('rows_short') else ''
            print(f"  {d['account']}: 台账 {d['rows']} 行 / scrape {d['scrape_items']}{extra}{warn}")
    print(f'verify: {level}')
    return 1 if level == 'ERROR' else 0

def main() -> int:
    ap = argparse.ArgumentParser(description='日更后台账同步与校验工具')
    sub = ap.add_subparsers(dest='cmd', required=True)
    w = sub.add_parser('write', help='同步台账（委托 build_ledger_15col write）')
    w.add_argument('--accounts', nargs='*', default=None)
    e = sub.add_parser('export-trans', help='导出 N 列待译清单')
    e.add_argument('--accounts', nargs='*', default=None)
    a = sub.add_parser('apply-trans', help='写入 N 列译文')
    a.add_argument('json_file')
    v = sub.add_parser('verify', help='终验（API + 台账磁盘复读）')
    v.add_argument('--skip-api', action='store_true', help='跳过工作台 API 校验')
    v.add_argument('--accounts', nargs='*', default=None, help='仅核验指定账号')
    args = ap.parse_args()
    if hasattr(args, 'accounts') and args.accounts == []:
        print('[error] --accounts 需要至少一个账号，未执行')
        return 1
    if args.cmd == 'write':
        extra = ['--accounts', *args.accounts] if args.accounts else []
        return run_builder(['write'] + extra)
    if args.cmd == 'export-trans':
        extra = ['--accounts', *args.accounts] if args.accounts else []
        return run_builder(['export-trans'] + extra)
    if args.cmd == 'apply-trans':
        return run_builder(['apply-trans', args.json_file])
    if args.cmd == 'verify':
        return main_verify(args.skip_api, args.accounts)
    return 1
if __name__ == '__main__':
    sys.exit(main())
