"""scrape_product_links.py — 抓取各账号商品链接并回填台账 G 列
用法：
  1) 抓取（tiktok env）:  python scrape_product_links.py scrape
  2) 回填 G 列（webapp venv）: python scrape_product_links.py fill
来源优先级：视频挂车商品链接（products.json，vid 级）> 主页 bio 链接（账号级）。
幂等：product_links.json 已存在的账号跳过。
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import json
import re
import sys
from pathlib import Path
XLSX = Path(os.path.join(os.environ.__getitem__('WB_LEDGER_DIR'), 'tiktok_对标.xlsx'))
LEDGER_DIR = Path(os.environ.__getitem__('WB_LEDGER_DIR'))
PROXY = os.environ.get('WB_PROXY', '')

def fetch_bio_link(account: str) -> str:
    """抓主页解析 bioLink（带货链接多为 Linktree/独立站/店铺直链）。"""
    from tiktok_fetch import fetch_html
    html = fetch_html(f'https://www.tiktok.com/@{account}')
    m = re.search('<script id="__UNIVERSAL_DATA_FOR_REHYDRATION__"[^>]*>(.*?)</script>', html, re.S)
    if not m:
        raise RuntimeError('主页无内嵌数据（风控）')
    data = json.loads(m.group(1))
    ui = ((data.get('__DEFAULT_SCOPE__', {}).get('webapp.user-detail') or {}).get('userInfo') or {}).get('user', {})
    return ((ui.get('bioLink') or {}).get('url') or '').strip()

def scrape_all() -> int:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    accounts = [dict(a) for a in json.loads((Path(__file__).resolve().parent.parent / 'accounts.json').read_text(encoding='utf-8')).get('accounts', [])]
    results = []
    for acc_cfg in accounts:
        acc = acc_cfg['account']
        region = str((acc_cfg or {}).get('region') or '').lower()
        if region == 'eu':
            print(f'[skip] {acc}（欧洲账号，按规则不抓商品链接）')
            continue
        plf = LEDGER_DIR / acc / f'{acc}_product_links.json'
        if plf.exists():
            data = json.loads(plf.read_text(encoding='utf-8'))
            print(f'[skip] {acc}（已有 product_links.json）')
            results.append((acc, data))
            continue
        vid_links = {}
        pj = LEDGER_DIR / acc / f'{acc}_products.json'
        if pj.exists():
            try:
                pdata = json.loads(pj.read_text(encoding='utf-8'))
                for vid, prods in pdata.items():
                    if isinstance(prods, list) and prods:
                        pid = str((prods[0] or {}).get('product_id') or '')
                        if pid:
                            vid_links[vid] = f'https://shop.tiktok.com/view/product/{pid}'
            except Exception:
                pass
        try:
            bio = fetch_bio_link(acc)
        except Exception as e:
            print(f'[warn] {acc} bio 抓取失败: {e.__class__.__name__} {str(e)[:80]}')
            bio = ''
        data = {'bio_link': bio, 'videos': vid_links}
        plf.parent.mkdir(parents=True, exist_ok=True)
        plf.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
        results.append((acc, data))
        print(f"[scraped] {acc}: bio={bio or '（无）'} 挂车链接 {len(vid_links)} 条")
        import time
        time.sleep(2)
    (XLSX.parent / 'product_links_results.json').write_text(json.dumps(dict(results), ensure_ascii=False, indent=1), encoding='utf-8')
    return 0

def fill_g() -> int:
    import openpyxl
    results = json.loads((XLSX.parent / 'product_links_results.json').read_text(encoding='utf-8'))
    wb = openpyxl.load_workbook(XLSX)
    for sn in wb.sheetnames:
        if not sn.startswith('对标记录-'):
            continue
        acc = sn.replace('对标记录-', '')
        if acc not in results:
            print(f'[skip-fill] {acc}（无抓取结果）')
            continue
        data = results[acc]
        ws = wb[sn]
        filled = 0
        for r in range(2, ws.max_row + 1):
            s = str(ws.cell(r, 6).value or '')
            m = re.search('ID:\\s*(\\d+)', s) or re.search('/video/(\\d+)', s)
            vid = m.group(1) if m else ''
            link = data.get('videos', {}).get(vid) or data.get('bio_link') or ''
            if link:
                ws.cell(r, 7).value = link
                filled += 1
        print(f'[fill] {acc}: G 列回填 {filled} 行')
    wb.save(XLSX)
    print('saved:', XLSX)
    return 0
if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'scrape'
    if cmd == 'scrape':
        sys.exit(scrape_all())
    elif cmd == 'fill':
        sys.exit(fill_g())
    else:
        print(__doc__)
        sys.exit(1)
